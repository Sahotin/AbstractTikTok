"""JSON and streaming JSONL readers with record-type discovery."""

from __future__ import annotations

import json
import csv
import asyncio
import io
from dataclasses import dataclass
from pathlib import Path
from typing import Any, AsyncIterator, Literal, Sequence

import aiofiles


RecordKind = Literal["content", "comment", "author"]


@dataclass(frozen=True)
class InputSource:
    path: Path
    kind: RecordKind
    sheet_name: str | None = None


async def iter_records(path: Path, sheet_name: str | None = None) -> AsyncIterator[dict[str, Any]]:
    """Yield records from crawler JSON, JSONL, CSV, or Excel output.

    JSONL is fully streaming. Existing MediaCrawler JSON output is one JSON
    array, so the standard library must materialize that array; the CLI prints
    this limitation and JSONL remains the recommended large-file format.
    """

    suffix = path.suffix.lower()
    if suffix == ".jsonl":
        async with aiofiles.open(path, "r", encoding="utf-8-sig") as file:
            line_number = 0
            async for line in file:
                line_number += 1
                stripped = line.strip()
                if not stripped:
                    continue
                try:
                    record = json.loads(stripped)
                except json.JSONDecodeError as exc:
                    raise ValueError(f"Invalid JSONL at {path}:{line_number}: {exc}") from exc
                if not isinstance(record, dict):
                    raise ValueError(f"Expected an object at {path}:{line_number}")
                yield record
        return

    if suffix == ".json":
        async with aiofiles.open(path, "r", encoding="utf-8-sig") as file:
            payload = json.loads(await file.read())
        if isinstance(payload, dict):
            yield payload
            return
        if not isinstance(payload, list):
            raise ValueError(f"Expected a JSON object or array in {path}")
        for index, record in enumerate(payload, start=1):
            if not isinstance(record, dict):
                raise ValueError(f"Expected an object at {path} record {index}")
            yield record
        return

    if suffix == ".csv":
        async with aiofiles.open(path, "r", encoding="utf-8-sig", newline="") as file:
            content = await file.read()
        for record in csv.DictReader(io.StringIO(content)):
            yield dict(record)
        return

    if suffix == ".xlsx":
        if not sheet_name:
            raise ValueError(f"Excel input requires a sheet name: {path}")
        records = await asyncio.to_thread(_read_excel_sheet, path, sheet_name)
        for record in records:
            yield record
        return

    raise ValueError(f"Unsupported input format: {path.suffix}. Expected .json, .jsonl, .csv, or .xlsx")


def _read_excel_sheet(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    """Read one MediaCrawler workbook sheet off the event loop."""

    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Excel sheet not found: {sheet_name}")
        rows = workbook[sheet_name].iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            return []
        names = [str(value).strip() if value is not None else "" for value in headers]
        return [
            {name: value for name, value in zip(names, row) if name}
            for row in rows
            if any(value is not None and value != "" for value in row)
        ]
    finally:
        workbook.close()


def _kind_from_name(path: Path) -> RecordKind | None:
    name = path.name.lower()
    if "comment" in name:
        return "comment"
    if "content" in name or "aweme" in name or "video" in name:
        return "content"
    if "creator" in name or "author" in name:
        return "author"
    return None


def _kind_from_record(record: dict[str, Any]) -> RecordKind:
    if "comment_id" in record:
        return "comment"
    if "aweme_id" in record:
        return "content"
    if "user_id" in record:
        return "author"
    raise ValueError("Cannot infer record type; expected comment_id, aweme_id, or user_id")


async def discover_input_sources(inputs: Sequence[Path | str]) -> list[InputSource]:
    """Expand files/directories, validate formats, and infer record types."""

    paths: set[Path] = set()
    for raw_input in inputs:
        path = Path(raw_input).expanduser().resolve()
        if not path.exists():
            raise FileNotFoundError(f"Input does not exist: {path}")
        if path.is_dir():
            paths.update(
                candidate
                for candidate in path.rglob("*")
                if candidate.suffix.lower() in {".json", ".jsonl", ".csv", ".xlsx"}
            )
        elif path.suffix.lower() in {".json", ".jsonl", ".csv", ".xlsx"}:
            paths.add(path)
        else:
            raise ValueError(f"Unsupported input format: {path}")

    if not paths:
        raise ValueError("No supported JSON, JSONL, CSV, or Excel input files were found")

    sources: list[InputSource] = []
    for path in sorted(paths):
        if path.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook

            workbook = await asyncio.to_thread(load_workbook, path, read_only=True, data_only=True)
            try:
                sheet_map = {"creators": "author", "contents": "content", "comments": "comment"}
                for sheet_name in workbook.sheetnames:
                    kind = sheet_map.get(sheet_name.strip().lower())
                    if kind:
                        sources.append(InputSource(path=path, kind=kind, sheet_name=sheet_name))
            finally:
                workbook.close()
            continue
        kind = _kind_from_name(path)
        if kind is None:
            first_record = None
            async for record in iter_records(path):
                first_record = record
                break
            if first_record is None:
                raise ValueError(f"Input file is empty: {path}")
            kind = _kind_from_record(first_record)
        sources.append(InputSource(path=path, kind=kind))

    if not sources:
        raise ValueError("No content, comment, or creator records were found in the selected inputs")

    order = {"author": 0, "content": 1, "comment": 2}
    return sorted(sources, key=lambda source: (order[source.kind], str(source.path)))
