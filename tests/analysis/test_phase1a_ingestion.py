from __future__ import annotations

import json
from datetime import timezone
from pathlib import Path
from uuid import uuid4

import pytest

from analysis.io import iter_records
from analysis.normalization import BilibiliNormalizer, DouyinNormalizer
from analysis.preprocessing import parse_datetime, validate_comment
from analysis.repositories import AnalysisRepository
from analysis.services import IngestionService


FIXTURES = Path(__file__).parent.parent / "fixtures"


def _load_content() -> dict:
    return json.loads((FIXTURES / "douyin_contents.json").read_text(encoding="utf-8"))[0]


def _load_comments() -> list[dict]:
    return [json.loads(line) for line in (FIXTURES / "douyin_comments.jsonl").read_text(encoding="utf-8").splitlines() if line]


def test_real_shape_douyin_comment_normalization() -> None:
    raw = _load_comments()[0]
    result = DouyinNormalizer().normalize_comment(raw, uuid4())

    assert result.native_comment_id == raw["comment_id"]
    assert result.native_content_id == raw["aweme_id"]
    assert result.text == raw["content"]
    assert result.like_count == 38
    assert result.parent_comment_id is None
    assert result.depth == 0
    assert result.language == "zh"
    assert len(result.text_hash) == 64


def test_real_shape_douyin_content_normalization() -> None:
    raw = _load_content()
    result = DouyinNormalizer().normalize_content(raw, uuid4())

    assert result.native_content_id == raw["aweme_id"]
    assert result.content_type == "video"
    assert result.like_count == 30744
    assert result.favorite_count == 8215
    assert result.view_count is None
    assert result.url == raw["aweme_url"]
    assert result.published_at is not None
    assert result.published_at.tzinfo == timezone.utc


def test_seconds_timestamp_is_parsed_as_utc() -> None:
    result = parse_datetime(1700000000)
    assert result is not None
    assert result.isoformat() == "2023-11-14T22:13:20+00:00"


def test_milliseconds_timestamp_is_parsed_as_utc() -> None:
    result = parse_datetime(1700000000000)
    assert result is not None
    assert result.isoformat() == "2023-11-14T22:13:20+00:00"


def test_missing_and_invalid_timestamps_are_explicit() -> None:
    assert parse_datetime(None) is None
    with pytest.raises(ValueError, match="Invalid datetime"):
        parse_datetime("not-a-time")


def test_second_level_comment_keeps_parent_root_and_depth() -> None:
    raw = _load_comments()[1]
    result = DouyinNormalizer().normalize_comment(raw, uuid4())

    assert result.parent_comment_id == "7634140513210762047"
    assert result.root_comment_id == "7634140513210762047"
    assert result.depth == 1


def test_invalid_comment_is_reported() -> None:
    raw = json.loads((FIXTURES / "douyin_invalid_comments.json").read_text(encoding="utf-8"))[0]
    result = DouyinNormalizer().normalize_comment(raw, uuid4())
    issues = validate_comment(result, {raw["aweme_id"]})

    assert any(issue.severity == "error" and "text is empty" in issue.message for issue in issues)


@pytest.mark.asyncio
async def test_jsonl_reader_yields_records_line_by_line(tmp_path: Path) -> None:
    path = tmp_path / "stream_comments.jsonl"
    path.write_text('{"comment_id":"1"}\n{"comment_id":"2"}\n', encoding="utf-8")

    records = []
    async for record in iter_records(path):
        records.append(record)

    assert records == [{"comment_id": "1"}, {"comment_id": "2"}]


@pytest.mark.asyncio
async def test_repeated_import_is_idempotent(tmp_path: Path) -> None:
    repository = AnalysisRepository(tmp_path / "analysis.db")
    service = IngestionService(repository, batch_size=1)
    inputs = [FIXTURES / "douyin_contents.json", FIXTURES / "douyin_comments.jsonl"]
    try:
        first = await service.ingest(inputs, platform="douyin")
        second = await service.ingest(inputs, platform="douyin")
    finally:
        await repository.close()

    assert first.run_id == second.run_id
    assert first.contents_inserted == 1
    assert first.comments_inserted == 2
    assert first.database_content_count == 1
    assert first.database_comment_count == 2
    assert first.first_level_comment_count == 1
    assert first.second_level_comment_count == 1
    assert second.contents_inserted == 0
    assert second.comments_inserted == 0
    assert second.content_duplicates == 1
    assert second.comment_duplicates == 2
    assert second.database_content_count == 1
    assert second.database_comment_count == 2


@pytest.mark.asyncio
async def test_ingestion_reports_invalid_rows(tmp_path: Path) -> None:
    repository = AnalysisRepository(tmp_path / "analysis.db")
    service = IngestionService(repository, batch_size=10)
    try:
        report = await service.ingest(
            [FIXTURES / "douyin_contents.json", FIXTURES / "douyin_invalid_comments.json"],
            platform="douyin",
        )
    finally:
        await repository.close()

    assert report.invalid_count == 1
    assert report.comments_inserted == 0
    assert report.database_comment_count == 0
    assert any(issue.entity_type == "comment" for issue in report.issues)


def test_real_shape_bilibili_comment_normalization() -> None:
    raw = {
        "comment_id": "9911",
        "parent_comment_id": "0",
        "video_id": "314480159984",
        "content": "这个观点值得讨论",
        "user_id": "42",
        "nickname": "测试用户",
        "create_time": 1_700_000_000,
        "like_count": "12",
        "sub_comment_count": "3",
        "last_modify_ts": 1_700_000_100,
    }

    result = BilibiliNormalizer().normalize_comment(raw, uuid4())

    assert result.platform.value == "bilibili"
    assert result.native_content_id == "314480159984"
    assert result.text == "这个观点值得讨论"
    assert result.like_count == 12
    assert result.depth == 0


@pytest.mark.asyncio
async def test_bilibili_excel_ingestion(tmp_path: Path) -> None:
    from openpyxl import Workbook

    workbook_path = tmp_path / "bili_detail.xlsx"
    workbook = Workbook()
    contents = workbook.active
    contents.title = "Contents"
    contents.append([
        "video_id", "video_type", "title", "desc", "create_time", "user_id",
        "nickname", "liked_count", "video_comment", "video_url", "last_modify_ts",
    ])
    contents.append([
        "314480159984", "video", "测试视频", "视频简介", 1_700_000_000, "42",
        "UP主", "100", "1", "https://www.bilibili.com/video/av314480159984", 1_700_000_100,
    ])
    comments = workbook.create_sheet("Comments")
    comments.append([
        "comment_id", "parent_comment_id", "create_time", "video_id", "content",
        "user_id", "nickname", "sub_comment_count", "like_count", "last_modify_ts",
    ])
    comments.append([
        "9911", "0", 1_700_000_010, "314480159984", "这个观点值得讨论",
        "99", "评论用户", "0", "12", 1_700_000_100,
    ])
    creators = workbook.create_sheet("Creators")
    creators.append(["user_id", "nickname", "total_fans", "last_modify_ts"])
    creators.append(["42", "UP主", "500", 1_700_000_100])
    workbook.save(workbook_path)

    repository = AnalysisRepository(tmp_path / "analysis.db")
    service = IngestionService(repository, batch_size=10)
    try:
        report = await service.ingest([workbook_path], platform="bilibili", crawler_type="detail")
        comments_page = await repository.get_comments_by_ids([])
    finally:
        await repository.close()

    assert report.contents_inserted == 1
    assert report.comments_inserted == 1
    assert report.authors_inserted >= 1
    assert report.database_comment_count == 1
    assert comments_page == []
